import json
from datetime import datetime, timezone, timedelta
from pathlib import Path
from openpyxl import load_workbook

here=Path(__file__).resolve().parent
evidence=here.parent/'evidence'
r=json.loads((evidence/'job_registry.json').read_text(encoding='utf-8'))
jobs={j['job_id']:j for j in r['jobs']}
stamp=datetime.now(timezone(timedelta(hours=8))).isoformat()
ids=['4339778','4339779']
new=[]
for jid in ids:
    job=jobs[jid]; spec=job['submission']
    mappings=spec['measurement_baselines']
    assert len(mappings)==12 and sum(v is not None for v in mappings.values())==4
    for name,base in mappings.items():
        if base:
            assert name.replace('_probe_','_anchor_')==base
            assert base in mappings and mappings[base] is None
    new.append({'job_id':jid,'state':job['state'],'host':job.get('exec_host'),'order':spec['environment']['CROSSOVER']})
data={'updated_at_hkt':stamp,'phase':'same-node balanced crossover submitted','jobs':new,
      'design':'A=PP anchor→PP probe→TG; B=PP anchor→TG→PP probe. ABBA and BAAB each contain four blocks. Every call retains r10. Only compare probe/anchor within its own block before combining evidence.',
      'next':'检查阶段和固定PPL，失败则定位日志。完成后验证24个测速进程各10样本、全量SHA检查；分别计算各段probe/anchor PP与CV，先比较同节点A/B并保留时间位置和跨段残留效应，再交叉核对另一顺序。对齐每TID stat/schedstat/status/wchan与GGUF驻留页、NUMA、主机计数；schedstat为0或权限不足时明确不可解释。只诊断不计Score，不改变正式参数，不删慢样本。',
      'shared_ppl':'each job has one fresh fixed PPL; all 12 diagnostic commands within that job share it; ±5%',
      'binary':'unchanged hash-verified Intel Classic production build',
      'excel':'D:/compete/Qwen38_CPU_实验日志.xlsx'}
(here/'next_actions.json').write_text(json.dumps(data,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
prior=here.parent/'order_diagnostic/next_actions.json'
p=json.loads(prior.read_text(encoding='utf-8'))
p['updated_at_hkt']=stamp;p['latest_followup']=str(here/'next_actions.json');p['next']=data['next']
prior.write_text(json.dumps(p,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
w=load_workbook('D:/compete/Qwen38_CPU_实验日志.xlsx',read_only=True)
rows={str(row[0]):row for row in w['作业日志'].iter_rows(min_row=2,values_only=True)}
for jid in ids:
    row=rows[jid]
    assert all(row[i] is not None for i in range(6))
    assert row[6] in ['PEND','RUN']
    assert all(row[i] is None for i in [14,15,16])
for row in w['运行结果'].iter_rows(min_row=2,values_only=True):
    if str(row[0]) in ids: assert row[10] is None
v={'checked_at_hkt':stamp,'job_ids':ids,'job_count':len(rows),'raw_sample_count':w['十次原始样本'].max_row-1,'all_registered':True,'within_block_baseline_mappings_verified':True,'diagnostic_score_blank':True}
w.close()
(here/'log_validation.json').write_text(json.dumps(v,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps({'jobs':new,'validation':v},ensure_ascii=False))
