import json
from datetime import datetime, timedelta, timezone
from pathlib import Path

from openpyxl import load_workbook

root=Path(__file__).resolve().parent
work=Path('D:/compete')
registry=json.loads((root.parent/'evidence/job_registry.json').read_text(encoding='utf-8'))
ids=['4340916','4340917']
jobs=[j for j in registry['jobs'] if j['job_id'] in ids]
assert len(jobs)==2
stamp=datetime.now(timezone(timedelta(hours=8))).isoformat()
actions={
    'updated_at_hkt':stamp,'phase':'HC-post v2 explicit AVX2/FMA with initial rounding barrier submitted',
    'jobs':jobs,
    'hypothesis':'For actual hc=4, explicit AVX2 chunks and scalar tails first round x*post behind a register compiler barrier, then apply four ordered FMA operations exactly as in baseline assembly; original flat per-thread partition and generic fallback remain. This differs from the previous generic HC row-segment rewrite and does not include HC-pre/NUMA changes.',
    'gates':'ICC build and quantize checks; actual-library 972-case graph differential with 1/3/38 threads and 3 repeats, double oracle and cross-build bit identity; 10 alternating operator trials; then each A/B case independently runs fixed PPL ±5% and full PP10/TG10.',
    'orders':{'4340916':'ABBA','4340917':'BAAB'},
    'interpretation':'Pending full results; operator speedup alone does not prove model gain. Keep all samples and do not replace production. Analyze within-node B1/A1 and B2/A2 plus both order directions.',
    'next':'Check build/differential gate. If failed diagnose and log before any repair/resubmission. If passed continue the already queued full PPL and balanced model tests. No duplicate submissions.',
    'excel':str(work/'Qwen38_CPU_实验日志.xlsx')}
(root/'next_actions.json').write_text(json.dumps(actions,ensure_ascii=False,indent=2),encoding='utf-8')
previous=root.parent/'barrier_tree/next_actions.json'
old=json.loads(previous.read_text(encoding='utf-8'))
old['latest_followup']={'updated_at_hkt':stamp,'path':str(root/'next_actions.json'),'jobs':ids,'purpose':'HC-post v2 repairs initial multiply contraction; v1 differential failures retained, previous runtime variants unadopted'}
previous.write_text(json.dumps(old,ensure_ascii=False,indent=2),encoding='utf-8')
book=load_workbook(work/'Qwen38_CPU_实验日志.xlsx',read_only=True,data_only=True)
sheet=book['作业日志']
headers=[c.value for c in next(sheet.iter_rows())]
rows=[dict(zip(headers,row)) for row in sheet.iter_rows(min_row=2,values_only=True) if str(row[0]) in ids]
assert len(rows)==2
for r in rows:
    assert r[headers[1]] is not None
    assert 'HC-post' in str(r)
validation={'updated_at_hkt':stamp,'jobs':len(registry['jobs']),'rows':rows,'total_samples':book['十次原始样本'].max_row-1}
(root/'excel_validation.json').write_text(json.dumps(validation,ensure_ascii=False,indent=2,default=str),encoding='utf-8')
book.close()
print(json.dumps({'new_job_rows':len(rows),'jobs':len(registry['jobs']),'total_samples':validation['total_samples'],'status':'PASS'},ensure_ascii=False))
