import json
import math
import statistics
import unicodedata
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter
from ppl_policy5 import add_policy_review

WORK = Path('D:/compete')
EVIDENCE = WORK/'.codex-deploy/optimization/evidence'
SOURCE = EVIDENCE/'job_registry.json'
TARGET = WORK/'Qwen38_CPU_实验日志.xlsx'
ROOT = '/work/apac2602/qwen38-cpu/optimization/20261018-v1'
BASE = '/work/apac2602/qwen38-cpu'
REV = 'd81235049384534c167caea52b85a694f6103d14'
PPL_CONFIG = 'PPL: t/tb=38,c=512,chunks=2,b=512,ub=256,KV=f16,ngl=0,NUMA=distribute,no-warmup,timeout=590s'
PERF_CONFIG = 'PP: p2048/n0; TG: p0/n128/d2048; t38,b2048,ub512,KV=f16,ngl0,r10,NUMA=distribute'
ICC = 'Intel Classic 2021.9 (icc/icpc), O3/precise/no-ftz/CORE-AVX512/qsimd-honor-fp-model'
DEFAULT = ICC+'; OpenMP=ON,BLAS=OFF,bind=close,places=cores,OMP_DYNAMIC=FALSE'

# Descriptions supplement machine-collected scheduler/configuration evidence.
CATALOG = {
'4336625': ('HC/NUMA单测','HC/NUMA候选单元测试；原计划随后构建及PPL/AB', 'Intel Classic icc/icpc，O3/precise/no-ftz/CORE-AVX512/qsimd-honor-fp-model；gnu++17 hc_test、gnu11 numa_test；单节点40 slots；GGML_NUMA_PLACE_VERBOSE=1,GGML_NUMA_PIN_CORES=0；单测失败后未执行38线程微基准、完整模型构建、PPL或PP/TG', '无', '按账户LSF原始输出与测试日志补录，非本轮两个PGO重提。HC 7175 cases/0 failures；NUMA测试2156页、2148页在目标node0、8页errors，FAIL: pages not on node 0；退出码1。保留失败，不计算分数。'),
'4335554': ('部署/构建/PPL','首次离线部署并用 GCC 12.2 构建参考版本','GCC 12.2,Release,GGML_NATIVE=ON,OpenMP=ON,BLAS=OFF; '+PPL_CONFIG,'官方 PPL 2.7696','初始历史作业；早于用户明确要求 Intel。仅做部署与 PPL，无完整性能分数。'),
'4335568': ('构建/PPL','改用 Intel LLVM ICX/ICPX，严格浮点模式','IntelLLVM 2023.1; -fp-model=strict -ffp-contract=off; OpenMP=ON,BLAS=OFF; '+PPL_CONFIG,'官方 PPL 2.7696','PPL 失败，停止 PP/TG。'),
'4335584': ('构建/PPL','ICX 数值对齐复测','IntelLLVM 2023.1; OpenMP=ON,BLAS=OFF; '+PPL_CONFIG+'; 历史附加浮点参数未完整留存','官方 PPL 2.7696','PPL 失败。历史额外参数不按当前配置反推。'),
'4335605': ('构建/PPL','ICX 改用 precise/FMA 对齐并显式链接系统 libm','IntelLLVM 2023.1; -fp-model=precise -ffp-contract=fast; libm.so.6; '+PPL_CONFIG,'官方 PPL 2.7696','PPL 仍失败，未测速。'),
'4335636': ('部署/构建/PPL','改用 Intel Classic，构建四个目标并验证部署',DEFAULT+'; '+PPL_CONFIG,'官方 PPL 2.7696','PPL PASS；确认为 READY 部署。'),
'4335664': ('Profiling 探测','探测 VTune 驱动/硬件采样/带宽收集可用性','VTune 2023.1; HPC bandwidth / HW hotspots / SW fallback 探测; 40 slots','无','硬件 hotspots 可用；带宽驱动不可用。'),
'4335666': ('Profiling','采集 PP/TG 硬件热点、线程与 CPU 分布',DEFAULT+'; '+PERF_CONFIG+'; VTune HW hotspots,10ms,采样窗口在测量阶段恢复','无','带采样开销，仅用于定位，不计正式 Score。PP barrier30.79%，TG Q8 dot54.49%。'),
'4335768': ('完整基线','Intel 原版无 VTune 完整基线（r13n16）',DEFAULT+'; '+PPL_CONFIG+'; '+PERF_CONFIG,'无','Intel 性能基线之一；跨节点比较须注明。'),
'4335769': ('完整基线','Intel 原版完整基线；从繁忙 r13n16 改投同配置 r13n02',DEFAULT+'; '+PPL_CONFIG+'; '+PERF_CONFIG,'无','发生排队后调整目标节点；保留完整十次样本。'),
'4335817': ('算子/微基准','Q8 顺序展开 u2/u4 与独立累加器 a2/a4 筛选',ICC+'; nb64/128/256/512; resident+256MiB stream; 5轮轮换; 1392数值用例/实现','Intel 原始 Q8 dot（同进程/同节点）','u2/u4 逐位一致，收益微小；a2/a4 改变舍入且退化，不采用。'),
'4335828': ('算子差分','tinyBLAS 数据复用补丁与原版差分及汇编对比',ICC+'; Q8_0 GEMM; m/n=4,5,7,8,15,16,31,64;k块=1,3,16,128;任务=1,3,38; padded strides','Intel 原始 tinyBLAS（同节点）','768 用例 PASS；主热点 2000 字节机器码与原版完全相同，不认定收益。'),
'4335832': ('PPL 筛选','Q8 四独立累加器完整 PPL 门槛测试',DEFAULT+'; q8-a4.patch; '+PPL_CONFIG,'官方 PPL 2.7696','PPL 通过；微基准退化，未继续完整 PP/TG。'),
'4335834': ('节点诊断','记录节点/线程计算与等待、tinyBLAS 形状',DEFAULT+'; GGML_CPU_NODE_STATS/SHAPE_STATS; PP warmup前4图; TG skip5后4图; '+PPL_CONFIG+'; '+PERF_CONFIG,'无','诊断有开销，不计分；发现 CONCAT ne2=1 时单线程拷贝。'),
'4335851': ('算子差分','CONCAT 实际动态库逐位差分',ICC+'; F32/I32,4轴,padded rows/transpose fallback,1/3/38 tasks; 顺序模拟任务','Intel 原始 CONCAT（实际库）','216 用例 PASS；无独立 PP/TG 分数。'),
'4335857': ('同节点完整对照','同一节点依次测试 baseline、native、concat',DEFAULT+'; 各版本独立 PPL→PP10→TG10; native:OpenMP=OFF; concat:按行拷贝; '+PPL_CONFIG+'; '+PERF_CONFIG,'作业内 baseline（同节点、顺序对照）','明细见“运行结果”。未做 ABBA，仍有顺序/缓存影响。'),
'4335859': ('节点诊断','跳过预热后记录正式 PP 测量窗口四个图',DEFAULT+'; 诊断构建; PP skip4,record4; '+PERF_CONFIG,'无','同诊断产物 PPL 已在4335834通过；不计算 Score。CONCAT关键计算4.428649s，等待163.879419 CPU秒。'),
'4335863': ('算子/微基准','Intel pragma unroll 对 tinyBLAS RN 内循环筛选',ICC+'; 768差分; 38线程,5轮交替,每轮3次GEMM; 通用形状','Intel 原始 tinyBLAS（同节点）','逐位 PASS，收益不足1%，未进入模型候选。'),
'4335868': ('算子/微基准','tinyBLAS pragma unroll 复测真实热点形状',ICC+'; m10240/n512/k块80及m2560/n512/k块192; 38线程,5轮交替,3次/轮','Intel 原始 tinyBLAS（同节点）','768差分 PASS；热点微基准约1.05%/0.44%提升，未证明端到端收益。'),
'4335869': ('算子/微基准','Q8 dot 复测真实 k_blocks=80/192',ICC+'; nb80/192; resident+256MiB stream; 5轮轮换','Intel 原始 Q8 dot（同节点）','u2 流式中位约0.3%提升；独立累加器退化。'),
'4335870': ('同节点完整对照','同一节点依次测试 baseline、q8-u2',DEFAULT+'; 各版本独立 PPL→PP10→TG10; '+PPL_CONFIG+'; '+PERF_CONFIG,'作业内 baseline（同节点、顺序对照）','明细见“运行结果”；未做 ABBA。'),
'4335883': ('并发差分/微基准','CONCAT 差分改为真实并发，补测真实 PP 形状',ICC+'; OpenMP 38;216差分;列3/9+512,行10240;5轮交替,每轮10次copy','Intel 原始 CONCAT（同节点、同进程）','216并发用例 PASS；PP形状热缓存copy约8.78/8.94ms降至0.054/0.058ms。'),
'4335888': ('并发差分/微基准','CONCAT 同时补测 PP 与 decode 小形状',ICC+'; OpenMP38;216差分;列3/9+1/512,行10240;5轮交替,10次copy/轮','Intel 原始 CONCAT（同节点、同进程）','216并发用例及实际形状逐位PASS；小形状也未见copy退化。微基准不能代替端到端结果。'),
}

BUILDS={'4335816':'native','4335818':'q8-u2','4335819':'q8-a4','4335820':'pp-reuse','4335827':'diagnostic','4335837':'mkl','4335838':'ipo','4335850':'concat'}
VARIANTS={'4335829':'native','4335831':'q8-u2','4335833':'pp-reuse','4335842':'mkl','4335852':'concat'}
RUNTIME={'4335821':('bt0-passive','KMP_BLOCKTIME=0,OMP_WAIT_POLICY=PASSIVE'),'4335822':('bt1-active','KMP_BLOCKTIME=1,OMP_WAIT_POLICY=ACTIVE'),'4335823':('bt20-active','KMP_BLOCKTIME=20,OMP_WAIT_POLICY=ACTIVE'),'4335824':('bt200-active','KMP_BLOCKTIME=200,OMP_WAIT_POLICY=ACTIVE'),'4335825':('spread','OMP_PROC_BIND=spread'),'4335826':('interleave','numactl --interleave=all')}
for jid, variant in BUILDS.items():
    delta = {'native':'OpenMP=OFF','mkl':'OpenMP=ON,BLAS=ON,Intel10_64lp/oneMKL LP64 threaded','ipo':'OpenMP=ON,BLAS=OFF,-ipo','diagnostic':'NODE_STATS/SHAPE_STATS macros'}.get(variant,'OpenMP=ON,BLAS=OFF; '+variant+' patch')
    CATALOG[jid]=('候选构建','离线构建 '+variant+'，检查四目标、动态库及量化测试',ICC+'; '+delta+'; cmake -j40; CPU-only','无','IPO 编译器内部错误010101_3096/#10014，未生成可测速构建。' if variant=='ipo' else '构建及 test-quantize-fns PASS；构建本身无性能分数。')
for jid, variant in VARIANTS.items():
    override = {'native':'; 覆盖 OpenMP=OFF','mkl':'; 覆盖 BLAS=ON,Intel10_64lp/oneMKL LP64 threaded'}.get(variant,'')
    CATALOG[jid]=('完整候选测速',variant+'：独立 PPL 门槛后完整 PP/TG',DEFAULT+'; variant='+variant+override+'; '+PPL_CONFIG+'; '+PERF_CONFIG,'4335769 Intel 原版（跨节点筛选）','mkl 数值失败，门槛自动停止测速。' if variant=='mkl' else '跨节点结果不直接证明优化收益；关键候选另做同节点对照。')
for jid,(policy,override) in RUNTIME.items():
    CATALOG[jid]=('运行时筛选','固定 Intel 原版代码，测试 '+policy,DEFAULT+'; PP/TG覆盖:'+override+'; PPL仍固定默认条件; '+PPL_CONFIG+'; '+PERF_CONFIG,'4335769 Intel 原版（跨节点筛选）','只作为运行时筛选；正式使用须按规则7.1申报。')
for jid,order in [('4336045','auto-default → auto-interleave → copy-interleave → copy-default'),('4336046','copy-default → copy-interleave → auto-interleave → auto-default')]:
    CATALOG[jid]=('同节点内存矩阵','四种模型加载/NUMA配置，顺序='+order,DEFAULT+'; 同一Intel原版产物; load-mode auto/none × 默认/numactl interleave=0,1; 每作业先独立执行一次固定PPL门槛; '+PPL_CONFIG+'; 每配置'+PERF_CONFIG+'; -v及每15秒轻量OS记录; 清除KMP/OMP等待及KMP/GOMP附加亲和性覆盖','作业内 auto-default（同节点）','4个case共享同一次固定PPL验证，非每case重跑；正序/逆序分属两个节点，仍有节点与顺序影响。开发自测不等于正式排名。')

registry = json.loads(SOURCE.read_text(encoding='utf-8'))
merged_jobs={job['job_id']:job for job in registry['jobs']}
for snapshot in sorted((EVIDENCE/'job_updates').glob('*.json')):
    update=json.loads(snapshot.read_text(encoding='utf-8'))
    for job in update['jobs']:
        existing=merged_jobs.get(job['job_id'])
        if existing and existing['state'] in ['DONE','EXIT'] and job['state'] in ['RUN','PEND']:
            continue
        terminal_update=job['state'] in ['DONE','EXIT'] and existing and existing['state'] in ['RUN','PEND']
        if existing is None or update['updated_at']>=registry['updated_at'] or terminal_update:
            job['evidence'].append({'path':str(snapshot),'text':'本行由PGO限定作业增量采集更新；其余历史行保留。采集时间='+update['updated_at']})
            merged=dict(existing or {})
            merged.update(job)
            merged_jobs[job['job_id']]=merged
    registry['updated_at']=max(registry['updated_at'],update['updated_at'])
registry['jobs']=sorted(merged_jobs.values(),key=lambda job:int(job['job_id']))
jobs = registry['jobs']
# Audited descriptive corrections preserve raw submissions and runtime configuration.
for correction_file in sorted(EVIDENCE.glob('submission_metadata_corrections*.json')):
    correction=json.loads(correction_file.read_text(encoding='utf-8'))
    for job in jobs:
        change=correction['corrections'].get(job['job_id'])
        if change:
            job['submission'].update(change['metadata'])
            job['evidence'].append({'path':str(correction_file),'text':json.dumps({'corrected_at':correction['corrected_at'],'reason':correction['reason'],'change':change},ensure_ascii=False)})
for job in jobs:
    spec = job.get('submission')
    if job['job_id'] not in CATALOG and spec:
        CATALOG[job['job_id']] = (spec['kind'],spec['purpose'],spec['config'],spec['baseline'],spec.get('note','已登记新作业；结果随LSF与原始JSON更新。'))
    if job['job_id'] not in CATALOG and any(path.endswith('.user-hc-numa.out') for path in job['logs']):
        CATALOG[job['job_id']] = ('账户并行HC/NUMA试验（补录）','HC/NUMA单测、构建与校验；由账户内其他工作提交，按LSF记录补录','已核对LSF资源：ssc-cpu,hg_6148,单节点40 slots；原始提交命令保留在LSF证据页。该并行项目的完整编译/运行配置尚待独立核对，未假定其与本次Q8/PGO配置相同','无','账户内并行项目，保留其作业号、时间、节点和状态；本次未修改或重提该项目，不据缺失指标计算分数。')
missing = {r['job_id'] for r in jobs}-CATALOG.keys()
if '4336681' in CATALOG:
    kind,purpose,config,base_label,note=CATALOG['4336681']
    CATALOG['4336681']=(kind,purpose,config+'；性能口径补充：摘取版用_cvtsh_ss转换FP16，原库用查表，helper路径不一致',base_label,note+' 204154数值用例PASS有效；此摘取版速度对照存在helper混杂，负向样本保留，须以真实完整库对照判断性能。')
if '4336672' in CATALOG:
    kind,purpose,config,base_label,note=CATALOG['4336672']
    CATALOG['4336672']=(kind,purpose,config,base_label,'内核编译完成后，旧版objdump不支持--disassemble=SYMBOL导致工具检查退出；未执行差分或微基准。已兼容重提4336681，保留失败。')
if missing:
    raise ValueError('Missing explicit job descriptions: '+str(missing))
manual = {}
if TARGET.exists():
    previous = load_workbook(TARGET)
    if '作业日志' in previous:
        ws = previous['作业日志']
        headers = [c.value for c in ws[1]]
        if '人工备注（保留）' in headers:
            col = headers.index('人工备注（保留）')+1
            manual = {str(ws.cell(r,1).value):ws.cell(r,col).value for r in range(2,ws.max_row+1)}
    previous.close()

book = Workbook()
info = book.active
info.title='说明'
info.append(['项目','内容'])
for row in [
('更新时间（香港）',registry['updated_at']),('作业数',len(jobs)),('LSF状态汇总','；'.join(f"{state}: {sum(j['state']==state for j in jobs)}" for state in ['DONE','EXIT','RUN','PEND'])),('覆盖范围','本任务初始部署、ICX/ICC构建、profiling、优化、失败及排队/运行作业；每个LSF作业一行。'),
('提交时间来源','已完成：LSF标准输出末尾 submitted 时间；运行中：bjobs -l；时区固定 UTC+08:00，精确到秒。'),
('Baseline 口径','作业日志列出该实验的比较对象；不存在比较对象时填“无”。正确性固定基准与性能比较基线分列。'),
('正式基准','PPL 2.7696，允许±5%；PP 46.929024，TG 3.295524 tokens/s；自测Score=100*sqrt(PP/basePP*TG/baseTG)。'),
('Score 使用条件','PPL PASS且PP/TG各完整10次，并且没有profiling/节点诊断开销。仅为开发自测，不是组委会正式成绩。'),
('对照局限','跨节点只用于筛选；同节点顺序对照仍有缓存/先后影响。单算子微基准不等于完整模型收益。'),
('更新流程','每次新提交登记作业ID/目的/配置/baseline；collect_job_registry.py抓取LSF与原始结果，write_excel_log.py原地更新本表。人工备注列会保留。'),
('缺失值','没有结果的指标单元格留空，状态明确写运行中/未运行/失败；不得将其填0。'),
('源码版本',REV),('远端证据目录',ROOT),('原始快照',str(SOURCE)),
]: info.append(row)

log = book.create_sheet('作业日志')
log_headers=['作业ID','提交时间（香港）','运行配置','做了什么','结果是什么','相对 Baseline','LSF状态','实验类型','运行节点','开始时间（香港）','结束时间（香港）','排队分钟','PPL','PPL状态','PP tokens/s','TG tokens/s','自测Score','PP相对变化','TG相对变化','Score相对变化','结果口径/限制','PPL校验基准','证据路径','人工备注（保留）']
log.append(log_headers)
detail = book.create_sheet('运行结果')
detail.append(['作业ID','版本/结果目录','PPL','PPL状态','PP tokens/s','PP标准差','PP CV','TG tokens/s','TG标准差','TG CV','自测Score','相对 Baseline','基线PP','基线TG','PP相对变化','TG相对变化','Score相对变化','比较类型','证据目录','算子规模/模式','原版中位ns/dot','候选中位ns/dot','算子配对中位速度变化','逐位差分','诊断返回码','诊断结果'])
kernel_samples = book.create_sheet('算子原始样本')
kernel_samples.append(['作业ID','候选','blocks','模式','轮次','函数','ns/dot','配对速度变化','口径','证据文件'])
samples = book.create_sheet('十次原始样本')
samples.append(['作业ID','版本','workload','样本序号','tokens/s','耗时ns','测量类型','证据文件'])
events = book.create_sheet('LSF与配置证据')
events.append(['作业ID','证据路径','原始记录（截取）'])
for item in registry.get('admission_evidence',[]):
    events.append(['无（提交准入检查）',item['path'],item['text'][:30000]])

all_measurements = {m['name']:m for j in jobs for m in j['measurements']}
jobs_by_id = {j['job_id']:j for j in jobs}
def scalar(m,key):
    a=m.get(key)
    return a[0] if isinstance(a,list) and a else None
def is_diag(m):
    return m['name'].startswith(('diagnostic','profile'))
def metrics(m):
    p=m.get('ppl') or {}
    pp=scalar(m,'pp'); tg=scalar(m,'tg')
    for key,x in [('pp',pp),('tg',tg)]:
        if x:
            assert (x['n_threads'],x['n_batch'],x['n_ubatch'],x['n_gpu_layers'],x['type_k'],x['type_v'])==(38,2048,512,0,'f16','f16')
            assert (x['n_prompt'],x['n_gen'],x['n_depth'])==((2048,0,0) if key=='pp' else (0,128,2048))
            assert len(x['samples_ns'])==len(x['samples_ts'])==10
    score = None
    if p.get('status')=='PASS' and pp and tg and not is_diag(m):
        assert len(pp['samples_ts'])==len(tg['samples_ts'])==10
        score=100*math.sqrt(pp['avg_ts']/46.929024*tg['avg_ts']/3.295524)
    return p,pp,tg,score
def baseline(m,jid):
    if m.get('baseline_name'):
        return None if m['name']==m['baseline_name'] else all_measurements.get(m['baseline_name'])
    if m['name'].startswith('confirm-'):
        spec=jobs_by_id[jid].get('submission',{})
        name=spec.get('baseline_case','confirm-baseline_{job_id}').format(job_id=jid)
        return None if m['name']==name else all_measurements.get(name)
    return all_measurements.get('llama_bench_results_4335769') if jid in VARIANTS or jid in RUNTIME else None
def dt(value):
    return datetime.fromisoformat(value).replace(tzinfo=None) if value else None

for job in jobs:
    jid=job['job_id']
    kind,purpose,config,base_label,note=CATALOG[jid]
    results=[]
    first=None
    for m in job['measurements']:
        p,pp,tg,score=metrics(m)
        bm=baseline(m,jid)
        bp,bpp,btg,bs=metrics(bm) if bm else ({},None,None,None)
        dpp=pp['avg_ts']/bpp['avg_ts']-1 if pp and bpp else None
        dtg=tg['avg_ts']/btg['avg_ts']-1 if tg and btg else None
        ds=score/bs-1 if score and bs else None
        comp='同节点顺序对照' if bm and m['name'].startswith(('confirm-','memory-','compiler-','q8-paired-')) else ('跨节点筛选' if bm else '无')
        if m.get('ppl_scope'): comp+='；'+m['ppl_scope']
        label=bm['name'] if bm else '无'
        if jid in ['4335554','4335568','4335584','4335605','4335636','4335832']:
            label='官方 PPL 2.7696'
        detail.append([jid,m['name'],p.get('ppl'),p.get('status','未运行'),pp['avg_ts'] if pp else None,pp['stddev_ts'] if pp else None,pp['stddev_ts']/pp['avg_ts'] if pp else None,tg['avg_ts'] if tg else None,tg['stddev_ts'] if tg else None,tg['stddev_ts']/tg['avg_ts'] if tg else None,score,label,bpp['avg_ts'] if bpp else None,btg['avg_ts'] if btg else None,dpp,dtg,ds,comp,m['path']])
        result=f"{m['name']}: PPL={p.get('ppl','待运行')}/{p.get('status','未运行')}"
        if pp: result+=f"; PP={pp['avg_ts']:.6f} (CV {100*pp['stddev_ts']/pp['avg_ts']:.2f}%)"
        if tg: result+=f"; TG={tg['avg_ts']:.6f} (CV {100*tg['stddev_ts']/tg['avg_ts']:.2f}%)"
        if score: result+=f'; Score={score:.3f}'
        results.append(result)
        if first is None:
            first=(p,pp,tg,score,dpp,dtg,ds)
        for phase,obj,filename in [('PP2048',pp,'pp2048.json'),('TG128@D2048',tg,'tg128_d2048.json')]:
            if not obj: continue
            for index,(speed,ns) in enumerate(zip(obj['samples_ts'],obj['samples_ns']),1):
                key='pp_file' if phase=='PP2048' else 'tg_file'
                samples.append([jid,m['name'],phase,index,speed,ns,'节点诊断（不计分）' if is_diag(m) else ('完整测速+15s轻量OS监控' if m.get('os_monitoring') else '无VTune的完整测速'),m.get(key,m['path']+'/'+filename)])
    if kind == '同节点完整对照' or len(job['measurements'])>1:
        first=None
    status_notes=[e['text'].strip() for e in job['evidence'] if e['path'].endswith(('/status','/job.status'))]
    summary='\n'.join(results) if results else note
    for evidence in job['evidence']:
        if evidence['path'].endswith('/relink-result.json'):
            diagnosis=json.loads(evidence['text'])
            diagnosis_lines=[]
            for item in diagnosis['results']:
                state='LINK PASS' if item['link_pass'] else 'LINK FAIL'
                diagnosis_lines.append(f"{item['variant']}: {state}, rc={item['returncode']}; 未执行产物、未加载模型")
                detail.append([jid,'relink-'+item['variant'],None,'未运行（仅链接）',None,None,None,None,None,None,None,'同作业relink-original（失败命令复现）',None,None,None,None,None,'同节点顺序original→no_as_needed→shared_intel；不执行；无PPL共享或Score',evidence['path'],None,None,None,None,None,item['returncode'],state+'；仅链接，运行/计数/PPL尚未验证'])
            summary+='\n'+'\n'.join(diagnosis_lines)
        if evidence['path'].endswith('/summary.json'):
            try:
                kernel_summary=json.loads(evidence['text'])
            except ValueError:
                continue
            if not isinstance(kernel_summary,dict):
                continue
            if kernel_summary.get('failure_reason'):
                summary+='\n'+kernel_summary['failure_reason']
            if kernel_summary.get('interpretation'):
                summary+='\n'+kernel_summary['interpretation']
            if 'kernel_times' in kernel_summary and 'checks' in kernel_summary:
                lines=[]
                for name,check in kernel_summary['checks'].items():
                    lines.append(f"{name}: {check['status']}, cases={check['cases']}, bit_mismatches={check['bit_mismatches']}")
                for result in kernel_summary['kernel_times']:
                    lines.append(f"{result.get('variant','')} nb{result['blocks']} {result['mode']}: 10轮配对中位速度变化={result['paired_median_speedup']:+.2%}")
                    variant=result.get('variant','avx512-ordered')
                    candidate=result.get(variant,{})
                    reference=result.get('baseline',{})
                    diagnostic_path=str(Path(evidence['path']).parent)
                    exact=kernel_summary['checks'].get(variant) or kernel_summary['checks'].get('library-correctness.json') or kernel_summary['checks'].get('correctness.json') or {}
                    detail.append([jid,variant,None,'未运行（算子诊断）',None,None,None,None,None,None,None,base_label,None,None,None,None,None,'同节点单核10轮AB/BA；配置依次执行；无PPL共享或Score',diagnostic_path,f"nb{result['blocks']} {result['mode']}",reference.get('ns_median'),candidate.get('ns_median'),result['paired_median_speedup'],f"{exact.get('status','未记录')} cases={exact.get('cases','未知')} mismatches={exact.get('bit_mismatches','未知')}"])
                    for name,stats in [('baseline',reference),(variant,candidate)]:
                        for index,ns in enumerate(stats.get('all_ns',[]),1):
                            kernel_samples.append([jid,variant,result['blocks'],result['mode'],index,name,ns,result['all_paired_speedups'][index-1],'原版与候选交替AB/BA；单核CPU0/membind0；非完整模型',evidence['path']])
                summary='\n'.join(results+lines)+'\n仅单核算子微基准，不是完整PP/TG，不计算Score。'
                if jid=='4336681':summary+='\n摘取版FP16 helper不同于实际库，速度比较有混杂；完整真实库对照作为主要性能依据。'
    if job['state'] in ['RUN','PEND']:
        summary+='\n当前阶段：'+('; '.join(status_notes) if status_notes else job['state'])
    elif job['state']=='EXIT' and status_notes:
        summary+='\n失败阶段：'+'; '.join(status_notes)
        if any('TRAIN_PP_UNSCORED' in s and 'exit=124' in s for s in status_notes):
            summary+='；PGO计数训练PP超时；control通过PPL，尚未生成PGO产物及完整PP/TG分数。'
        if jid in ['4336529','4336530']:
            summary+='；计算节点无git，源码检查尚未开始编译；已改为登录节点保存源码哈希、计算节点逐文件验证后重新提交。'
    lsf_args=job.get('submission',{}).get('lsf_args',[])
    if job['state']=='PEND' and '-w' in lsf_args:
        summary+='\n等待依赖：'+lsf_args[lsf_args.index('-w')+1]
    p,pp,tg,score,dpp,dtg,ds=first or ({},None,None,None,None,None,None)
    display_ppl_state = '多版本见明细' if kind=='同节点完整对照' or len(job['measurements'])>1 else p.get('status','待结果' if job['state'] in ['RUN','PEND'] else '未单独运行')
    if 'PGO' in kind and len(job['measurements'])==1 and job['measurements'][0]['name'].startswith('compiler-control_') and p.get('status')=='PASS':
        display_ppl_state='仅control PASS；PGO'+('待校验' if job['state'] in ['RUN','PEND'] else '未校验')
    start=dt(job.get('start_time')); submit=dt(job.get('submit_time')); finish=dt(job.get('finish_time'))
    queue=(start-submit).total_seconds()/60 if start and submit else None
    host=job.get('exec_host','尚未分配' if job['state']=='PEND' else '未记录')
    if host=='未记录' and jid=='4335857': host='40*r13n09'
    if host=='未记录' and jid=='4335870': host='40*r13n10'
    paths=job['logs']+[m['path'] for m in job['measurements']]
    paths+=sorted({m['ppl_gate_path'] for m in job['measurements'] if m.get('ppl_gate_path')})
    log.append([jid,submit,config,purpose,summary,base_label,job['state'],kind,host,start,finish,queue,p.get('ppl'),display_ppl_state,pp['avg_ts'] if pp else None,tg['avg_ts'] if tg else None,score,dpp,dtg,ds,note,2.7696 if job['measurements'] else '无','\n'.join(paths),manual.get(jid)])
    events.append([jid,'LSF completion summary',job.get('lsf_summary','运行中；由 bjobs -l 提取')[:30000]])
    if job.get('submission'):
        events.append([jid,'已登记提交参数',json.dumps(job['submission'],ensure_ascii=False,indent=2)])
    for item in job['evidence']:
        events.append([jid,item['path'],item['text'][:30000]])

# Profile data are retained separately and never included in normalized scores.
profile=WORK/'.codex-deploy/profiling/results_4335666'
for phase,sub in [('PP2048','pp2048'),('TG128@D2048','tg128_d2048')]:
    path=profile/sub/'bench.json'
    if path.exists():
        x=json.loads(path.read_text())[0]
        for index,(speed,ns) in enumerate(zip(x['samples_ts'],x['samples_ns']),1):
            samples.append(['4335666','VTune',phase,index,speed,ns,'VTune采样（不计分）',str(path)])

config_sheet=book.create_sheet('公共配置')
config_sheet.append(['配置项','固定值/来源'])
for row in [('资源','ssc-cpu; 单节点40 slots; 2xGold6148; 正式推理38线程'),('源码',REV),('Intel默认',DEFAULT),('完整浮点参数','-gcc-name/-gxx-name=GCC11.2; -fp-model precise -no-ftz -xCORE-AVX512 -qsimd-honor-fp-model -diag-disable=10441; Release O3 DNDEBUG'),('链接','-Wl,--no-as-needed,/lib64/libm.so.6,--as-needed'),('固定PPL',PPL_CONFIG),('固定PP/TG',PERF_CONFIG),('PPL门槛','[2.63112,2.90808]；模型路径、分片大小与WikiText哈希对照官方JSON'),('WikiText SHA256','173c87a53759e0201f33e0ccf978e510c2042d7f2cb78229d9a50d79b9e7dd08'),('模型','/work/ssc-zhaojh/llama_model/unsloth/Qwen3.8-Flash-Next-GGUF/UD-Q4_K_XL/Qwen3.8-Flash-Next-UD-Q4_K_XL-00001-of-00004.gguf'),('构建入口',ROOT+'/scripts/build_variant.lsf'),('评测入口',ROOT+'/scripts/benchmark_variant.lsf'),('同节点对照',ROOT+'/scripts/confirm_variants.lsf; confirm_q8.lsf'),('默认环境','OMP_NUM_THREADS=38; OMP_DYNAMIC=FALSE; OMP_PROC_BIND=close; OMP_PLACES=cores; OPENBLAS/MKL/BLIS_NUM_THREADS=38'),('运行时差异','每行配置显式记录覆盖项；性能策略变更不改PPL固定条件；没有新增额外推理线程')]:
    config_sheet.append(row)

if registry.get('rejected_submissions'):
    rejection_sheet=book.create_sheet('提交拒绝')
    rejection_sheet.append(['返回时间（香港）','做了什么','运行配置','相对 Baseline','提交返回码','LSF原始返回','重提来源作业','提交脚本','说明'])
    for submission in registry['rejected_submissions']:
        rejection_sheet.append([dt(submission['submitted_at_client']),submission['purpose'],submission['config'],submission['baseline'],submission['return_code'],submission['submission_output'],submission.get('retry_of','无'),submission['script'],'未取得LSF作业号，不计入作业数，不记为PEND/RUN；返回时间为客户端完成重试时刻。'])

add_policy_review(book, jobs, EVIDENCE)

for ws in book:
    ws.freeze_panes='B2' if ws.max_column>2 else 'A2'
    ws.sheet_view.showGridLines=False
    ws.row_dimensions[1].height=30
    for cell in ws[1]:
        cell.fill=PatternFill('solid',fgColor='17365D')
        cell.font=Font(name='Microsoft YaHei',color='FFFFFF',bold=True,size=11)
        cell.alignment=Alignment(vertical='center',wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if cell.data_type == 'f': cell.data_type = 's'
            cell.font=Font(name='Microsoft YaHei',size=10)
            cell.alignment=Alignment(vertical='top',wrap_text=True)
            if isinstance(cell.value,datetime): cell.number_format='yyyy-mm-dd hh:mm:ss'
    if ws.max_row>1:
        table=Table(displayName='T'+str(book.index(ws)),ref=ws.dimensions)
        table.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True)
        ws.add_table(table)
    ws.auto_filter.ref=ws.dimensions
    ws.print_options.horizontalCentered=False
    ws.sheet_properties.pageSetUpPr.fitToPage=True
    ws.page_setup.orientation='landscape'
    ws.page_setup.paperSize=ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
    ws.print_title_rows='1:1'
    for col in range(1,ws.max_column+1): ws.column_dimensions[get_column_letter(col)].width=20

for col,width in {'A':13,'B':23,'C':72,'D':40,'E':72,'F':37,'G':12,'H':20,'I':18,'J':23,'K':23,'L':13,'M':12,'N':14,'O':16,'P':16,'Q':15,'R':15,'S':15,'T':16,'U':65,'V':15,'W':75,'X':40}.items(): log.column_dimensions[col].width=width
for r in range(2,log.max_row+1):
    line_counts=[]
    for col in [3,4,5,6,21,23]:
        width=log.column_dimensions[get_column_letter(col)].width
        value=str(log.cell(r,col).value or '')
        line_counts.append(sum(max(1,math.ceil(sum(2 if unicodedata.east_asian_width(c) in 'WF' else 1 for c in line)/width)) for line in value.split('\n')))
    log.row_dimensions[r].height=min(320,max(80,14*(max(line_counts)+1)))
    for col in [12,13,15,16,17]: log.cell(r,col).number_format='0.000000' if col!=12 else '0.00'
    for col in [18,19,20]: log.cell(r,col).number_format='+0.00%;-0.00%;0.00%'
    state=log.cell(r,7).value
    log.cell(r,7).fill=PatternFill('solid',fgColor={'EXIT':'FFC7CE','RUN':'FFF2CC','PEND':'FFE699','DONE':'E2F0D9'}.get(state,'FFFFFF'))
    if log.cell(r,14).value=='FAIL': log.cell(r,14).fill=PatternFill('solid',fgColor='FFC7CE')
    if log.cell(r,14).value=='PASS': log.cell(r,14).fill=PatternFill('solid',fgColor='E2F0D9')
for r in range(2,detail.max_row+1):
    detail.row_dimensions[r].height=46
    for col in [3,5,6,8,9,11,13,14]: detail.cell(r,col).number_format='0.000000'
    for col in [7,10,15,16,17]: detail.cell(r,col).number_format='0.00%'
    detail.cell(r,23).number_format='+0.00%;-0.00%;0.00%'
detail.column_dimensions['B'].width=37;detail.column_dimensions['L'].width=45;detail.column_dimensions['S'].width=95
for ws in [info,config_sheet]:
    ws.column_dimensions['A'].width=29;ws.column_dimensions['B'].width=120
    for r in range(2,ws.max_row+1):ws.row_dimensions[r].height=44

if '提交拒绝' in book:
    rejection_sheet.column_dimensions['A'].width=23
    for column in ['B','C','D','F','H','I']:rejection_sheet.column_dimensions[column].width=65
    for r in range(2,rejection_sheet.max_row+1):rejection_sheet.row_dimensions[r].height=145
events.column_dimensions['B'].width=85;events.column_dimensions['C'].width=120
for r in range(2,events.max_row+1):events.row_dimensions[r].height=70
samples.column_dimensions['B'].width=38;samples.column_dimensions['H'].width=100
samples.column_dimensions['E'].width=18
for r in range(2,samples.max_row+1):samples.cell(r,5).number_format='0.000000'
book.active=book.sheetnames.index('作业日志')
book.properties.title='Qwen3.8 CPU 实验作业日志'
book.properties.description='LSF作业、提交时间、配置、改动、结果、Baseline；保留失败和完整十次样本。'
temp=TARGET.with_name(TARGET.stem+'.writing.xlsx')
book.save(temp)
book.close()
check=load_workbook(temp,data_only=True)
ids=[str(r[0]) for r in check['作业日志'].iter_rows(min_row=2,values_only=True)]
assert len(ids)==len(set(ids))==len(jobs)
assert set(ids)=={j['job_id'] for j in jobs}
for row in check['作业日志'].iter_rows(min_row=2,values_only=True):
    assert all(row[i] is not None for i in range(6)), row[:6]
    assert isinstance(row[1],datetime)
assert not any(c.data_type=='f' for ws in check for row in ws for c in row)
validation={'path':str(TARGET),'jobs':len(ids),'result_rows':check['运行结果'].max_row-1,'sample_rows':check['十次原始样本'].max_row-1,'updated_at':registry['updated_at']}
check.close()
temp.replace(TARGET)
print(json.dumps(validation,ensure_ascii=False))
