import json
import math
from datetime import datetime, timezone, timedelta
from pathlib import Path
from openpyxl import load_workbook

here=Path(__file__).resolve().parent
evidence=here.parent/'evidence'
registry=json.loads((evidence/'job_registry.json').read_text(encoding='utf-8'))
by_id={j['job_id']:j for j in registry['jobs']}
stamp=datetime.now(timezone(timedelta(hours=8))).isoformat()
target=evidence/'analysis_updates/pgo_mkl_20261019.json'
analysis=json.loads(target.read_text(encoding='utf-8'))
results={}
for jid in ['4338130','4338131']:
    job=by_id[jid]
    cases={m['name'].split('_')[0].rsplit('-',1)[1]:m for m in job['measurements']}
    a,b=cases['control'],cases['pgo']
    for m in [a,b]:
        assert m['ppl']['status']=='PASS'
        assert len(m['pp'][0]['samples_ts'])==len(m['tg'][0]['samples_ts'])==10
    pp=b['pp'][0]['avg_ts']/a['pp'][0]['avg_ts']-1
    tg=b['tg'][0]['avg_ts']/a['tg'][0]['avg_ts']-1
    geo=math.sqrt((1+pp)*(1+tg))-1
    results[jid]={'pp_change_percent':100*pp,'tg_change_percent':100*tg,'geomean_change_percent':100*geo}
    order='control→PGO' if jid=='4338130' else 'PGO→control'
    summary=(f'完整同节点对照已完成，{job["exec_host"]}，顺序{order}。两版本新PPL均2.7767 PASS，PP/TG各完整10次。'
             f'PGO相对本作业control：PP {pp:+.2%}，TG {tg:+.2%}，几何均值 {geo:+.2%}。'
             '正反序PGO的PP收益反转，且两轮后运行版本PP均明显较快；存在强烈顺序/缓存/节点或时段混杂，尚不能认定PGO有稳定收益。保留全部慢样本，暂不提升生产版本。开发自测非正式成绩。')
    analysis['jobs'][jid]={'summary':summary,'metrics':results[jid]}
analysis['updated_at_hkt']=stamp
target.write_text(json.dumps(analysis,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
new={str(j):by_id[str(j)] for j in range(4338277,4338282)}
next_actions={'updated_at_hkt':stamp,'phase':'five independent factor screens submitted',
              'jobs':[{'job_id':k,'state':v['state'],'host':v.get('exec_host'),'factor':v['submission']['environment']['FACTOR']} for k,v in new.items()],
              'pgo_paired_results':results,
              'next':'检查5组阶段状态；失败保留日志并按具体错误修复；完成后验证两版PPL及PP/TG全部10次，比较各作业同节点baseline。对有前景因素反序复测，避免将先后顺序效应误判为收益。Intel编译器、固定评测参数及PPL±5%不变。',
              'do_not_duplicate':['4338133','4338267','4338268'],
              'excel':'D:/compete/Qwen38_CPU_实验日志.xlsx'}
(here/'next_actions.json').write_text(json.dumps(next_actions,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
for name in ['pgo_hot_serial','compiler_pgo']:
    path=here.parent/name/'next_actions.json'
    previous=json.loads(path.read_text(encoding='utf-8'))
    previous['latest_followup']={'at':stamp,'completed_jobs':['4338130','4338131'],'conclusion':'正反序PP收益反转，不认定PGO稳定提升；5组其他变量已提交。','next_actions':str(here/'next_actions.json')}
    path.write_text(json.dumps(previous,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
book=load_workbook(Path('D:/compete/Qwen38_CPU_实验日志.xlsx'),read_only=True)
rows={str(row[0]):row for row in book['作业日志'].iter_rows(min_row=2,values_only=True)}
for jid in new:
    row=rows[jid]
    assert row[1] and row[2] and row[3] and row[5]
    assert row[6]=='RUN'
    assert all(row[i] is None for i in [12,14,15,16])
assert book['十次原始样本'].max_row-1>=1020
validation={'checked_at':stamp,'job_ids':list(new),'all_registered':True,'new_jobs_no_premature_metrics':True,'retained_raw_samples':book['十次原始样本'].max_row-1}
book.close()
(here/'log_validation.json').write_text(json.dumps(validation,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps({'pgo':results,'validation':validation},ensure_ascii=False))
