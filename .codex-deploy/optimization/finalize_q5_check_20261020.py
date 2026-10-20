import json
from pathlib import Path
from datetime import datetime, timezone, timedelta
from openpyxl import load_workbook

base = Path('D:/compete/.codex-deploy/optimization')
now = datetime.now(timezone(timedelta(hours=8))).isoformat()
for jid, tag in [('4340602', 'unroll'), ('4340624', 'combined')]:
    summary = json.loads((base / 'evidence' / (jid + '-summary.json')).read_text(encoding='utf-8'))
    review = summary['source_screening_review']
    path = base / ('q5_' + tag + '_20261020') / 'next_actions.json'
    value = json.loads(path.read_text(encoding='utf-8'))
    value.update(updated_hong_kong=now, last_known_status='DONE', latest_stage=review['phase'],
                 latest_stage_checked=now, review=review,
                 next=['Retain independent source and all raw model/kernel samples.',
                       'Do not promote: first/last original PP drift; full-model gain not established.',
                       'Combined candidate needs counterbalanced full-model replication before any performance claim.'])
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')

book = load_workbook(Path('D:/compete') / 'Qwen38_CPU_实验日志.xlsx', read_only=True, data_only=True)
for jid, expected in [('4340602', 100), ('4340624', 240)]:
    row = next(r for r in book['作业日志'].iter_rows(min_row=2, values_only=True) if str(r[0]) == jid)
    assert row[6] == 'DONE', row
    assert row[16] is None, 'Multi-case job must not be assigned one causal score'
    raw = [r for r in book['十次原始样本'].iter_rows(min_row=2, values_only=True) if str(r[0]) == jid]
    assert len(raw) == 60, (jid, len(raw))
    kernels = [r for r in book['算子原始样本'].iter_rows(min_row=2, values_only=True) if str(r[0]) == jid]
    assert len(kernels) == expected, (jid, len(kernels))
    print(json.dumps({'job_id': jid, 'state': row[6], 'model_samples': len(raw), 'kernel_samples': len(kernels)}, ensure_ascii=False))
book.close()
