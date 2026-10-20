import json
from datetime import datetime, timezone, timedelta
from pathlib import Path
from openpyxl import load_workbook

here=Path(__file__).resolve().parent
evidence=here.parent/'evidence'
registry=json.loads((evidence/'job_registry.json').read_text(encoding='utf-8'))
jobs={j['job_id']:j for j in registry['jobs']}
stamp=datetime.now(timezone(timedelta(hours=8))).isoformat()
ids=['4338534','4338535','4338536']
new=[]
for jid in ids:
    job=jobs[jid]
    assert job['submission']['script'].endswith('/20261019-order-diagnostic/order_job.lsf')
    new.append({'job_id':jid,'state':job['state'],'host':job.get('exec_host'),'mode':job['submission']['environment']['ORDER_MODE']})
data={'updated_at_hkt':stamp,'phase':'order-effect diagnostic and unsampled repeat submitted','jobs':new,
      'source_baseline':'4338278 hash-verified Intel original production binary; same-node first PP/TG are per-job baselines',
      'next':'检查阶段/PPL，诊断两作业均3次PP10与2次TG10，只变顺序。读各case pp/tg.observations.jsonl、.result.json、.stderr.log、.command.json，将缺页、read_bytes、GGUF NUMA分布和主机迁移计数按实际时间对齐；pp-chain与tg-between跨节点，不能仅用绝对速度作因果判断。AAA无采样三轮固定PPL/PP10/TG10，检查重复趋势，保留第一轮与所有慢样本。诊断绝不计算Score。',
      'fixed_requirements':'Intel Classic; 38 inference threads; fixed workloads and PPL ±5%; hg_6148 40slots; no OS settings changed',
      'excel':'D:/compete/Qwen38_CPU_实验日志.xlsx'}
(here/'next_actions.json').write_text(json.dumps(data,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
prior=here.parent/'factor_reverse/next_actions.json'
old=json.loads(prior.read_text(encoding='utf-8'))
old['phase']='all reverse jobs completed; order diagnostics submitted'
old['updated_at_hkt']=stamp
old['latest_followup']=str(here/'next_actions.json')
old['next']=data['next']
for job in old['jobs']: job['state']=jobs[job['job_id']]['state']
prior.write_text(json.dumps(old,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
w=load_workbook(Path('D:/compete/Qwen38_CPU_实验日志.xlsx'),read_only=True)
rows={str(r[0]):r for r in w['作业日志'].iter_rows(min_row=2,values_only=True)}
for jid in ids:
    row=rows[jid]
    assert all(row[i] is not None for i in range(6))
    assert row[6] in ['RUN','PEND']
    assert row[14] is None and row[15] is None and row[16] is None
assert rows['4338446'][6]=='DONE'
for row in w['运行结果'].iter_rows(min_row=2,values_only=True):
    if str(row[0]) in ids[:2]: assert row[10] is None
validation={'checked_at_hkt':stamp,'new_ids':ids,'all_registered':True,'diagnostic_unscored':True,'jobs':len(rows),'raw_samples':w['十次原始样本'].max_row-1}
w.close()
(here/'log_validation.json').write_text(json.dumps(validation,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps({'jobs':new,'validation':validation},ensure_ascii=False))
