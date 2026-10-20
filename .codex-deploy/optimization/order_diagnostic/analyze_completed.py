import json
from datetime import datetime, timezone, timedelta
from pathlib import Path
from openpyxl import load_workbook

here=Path(__file__).resolve().parent
evidence=here.parent/'evidence'
registry=json.loads((evidence/'job_registry.json').read_text(encoding='utf-8'))
jobs={j['job_id']:j for j in registry['jobs']}
telemetry=json.loads((here/'telemetry_summary.json').read_text(encoding='utf-8'))
stamp=datetime.now(timezone(timedelta(hours=8))).isoformat()
analysis={'updated_at_hkt':stamp,'source':'LSF completion summaries, fixed PPL and all benchmark samples, per-process read-only observations','jobs':{}}
for jid in ['4338534','4338535','4338536']:
    job=jobs[jid]
    assert job['state']=='DONE'
    vals=[]
    for m in job['measurements']:
        p=m.get('ppl')
        assert p and p['status']=='PASS' and p['ppl']==2.7767
        assert p['elapsed_seconds']<=590
        row={'name':m['name'],'ppl':p['ppl'],'ppl_seconds':p['elapsed_seconds']}
        for phase in ['pp','tg']:
            if m.get(phase):
                x=m[phase][0]
                assert len(x['samples_ns'])==len(x['samples_ts'])==10
                assert (x['n_threads'],x['n_batch'],x['n_ubatch'],x['type_k'],x['type_v'],x['n_gpu_layers'])==(38,2048,512,'f16','f16',0)
                assert (x['n_prompt'],x['n_gen'],x['n_depth'])==((2048,0,0) if phase=='pp' else (0,128,2048))
                row[phase]=x['avg_ts']; row[phase+'_cv_percent']=100*x['stddev_ts']/x['avg_ts']
        vals.append(row)
    assert len(vals)==3
    if jid=='4338534':
        summary='连续PP诊断完成，顺序PP1→PP2→PP3→TG1→TG2；三次PP均值31.096765→45.666116→28.507846 t/s，未稳定，第三次PP CV约29.6%；TG约2.67/2.62且有明显低速样本。全部5个测速进程major faults=0，采样read_bytes=0；GGUF驻留页数量在部分阶段明显变化，主线程上下文切换增加，主机NUMA迁移计数增长。不能仅凭主机计数或与另一节点的差异证明NUMA因果，也不能认定连续PP必然变慢。每30秒采样有开销，单次最高约0.85秒；不计Score，保留所有样本。'
    elif jid=='4338535':
        summary='穿插TG诊断完成，顺序PP1→TG1→PP2→TG2→PP3；PP31.907140→53.628610→53.646353 t/s，后两次接近，CV约0.17%/0.24%；TG3.007951/3.040360。全部5个测速进程major faults=0，采样read_bytes=0，后两次PP的GGUF页分布较稳定。TG穿插与稳定速度同时出现，但与连续PP作业不同节点/时段，尚不能认定TG是加速原因；主机NUMA迁移计数也非该进程独占。每30秒采样有开销，单次最高约1.19秒；不计Score。'
    else:
        increase=100*(vals[2]['pp']/vals[1]['pp']-1)
        summary=f'不采样三轮相同Intel原版完整固定评测均完成，各自独立PPL2.7767 PASS（313/72/92秒），每轮PP/TG各10次。PP31.341873→52.861537→53.370862 t/s；第三轮相对第二轮{increase:+.2f}%，后两轮接近。TG3.063875→3.026859→3.048581。复现相同配置先慢后快，支持顺序/运行状态混杂确实存在，不能作为代码优化收益；本节点后三轮中后两轮趋稳不能推广为所有节点都稳定。第一轮及全部慢样本保留，开发自测非正式成绩。'
    analysis['jobs'][jid]={'summary':summary,'values':vals,'telemetry_summary':telemetry.get(jid,[]),'evidence_files':[str(here/'completed_telemetry.json'),str(here/'telemetry_summary.json')]}
(evidence/'analysis_updates/order_diagnostic_20261020.json').write_text(json.dumps(analysis,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
(here/'completed_analysis.json').write_text(json.dumps(analysis,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
p=here/'next_actions.json'
next_actions=json.loads(p.read_text(encoding='utf-8'))
next_actions['updated_at_hkt']=stamp
next_actions['phase']='all three completed and analyzed'
for j in next_actions['jobs']: j['state']='DONE'
next_actions['next']='三作业已完成，不重提原作业。未经采样AAA与TG穿插诊断后两次PP均接近53t/s；连续PP在r13n10显著波动。全部诊断进程major faults=0且采样read_bytes=0，不支持简单磁盘读入解释。下一步优先同节点平衡顺序诊断，并关联GGUF驻留页变化、逐线程等待/调度、NUMA迁移；目前关联不足以确定因果。不得修改正式评测参数、挑选后两轮替代第一轮，或把预热效应称为优化收益。'
p.write_text(json.dumps(next_actions,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
w=load_workbook('D:/compete/Qwen38_CPU_实验日志.xlsx',read_only=True)
logs={str(r[0]):r for r in w['作业日志'].iter_rows(min_row=2,values_only=True)}
counts={jid:0 for jid in analysis['jobs']}
for row in w['十次原始样本'].iter_rows(min_row=2,values_only=True):
    if str(row[0]) in counts: counts[str(row[0])]+=1
assert counts=={'4338534':50,'4338535':50,'4338536':60},counts
for row in w['运行结果'].iter_rows(min_row=2,values_only=True):
    if str(row[0]) in ['4338534','4338535']: assert row[10] is None
for jid in counts: assert logs[jid][6]=='DONE'
validation={'checked_at_hkt':stamp,'sample_counts':counts,'diagnostic_scores_blank':True,'states_all_done':True,'total_raw_samples':w['十次原始样本'].max_row-1}
w.close()
(here/'completed_log_validation.json').write_text(json.dumps(validation,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps(validation))
